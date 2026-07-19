# -*- coding: utf-8 -*-
"""
Markdown → Word (.docx) dönüştürücü. Türkçe karakter destekli.

Desteklenen alt küme: # / ## / ### başlık, **kalın**, `kod`, | tablo |,
- / 1. liste, > alıntı, ``` kod bloğu, --- yatay çizgi.

Çalıştırma:
    python scripts/md_to_docx.py docs/GELISMIS_Sistem_Mimarisi_v3.md
    python scripts/md_to_docx.py <girdi.md> [cikti.docx]
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

ACCENT      = RGBColor(0x1F, 0x3B, 0x63)
GRI         = RGBColor(0x5A, 0x5A, 0x5A)
KOD_ARKA    = "F2F2F2"
BASLIK_ARKA = "D9E1F2"
ALINTI_ARKA = "FFF4E5"


def _shade(el, renk_hex):
    """Hücre veya paragraf gölgelendirmesi."""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), renk_hex)
    el.append(shd)


def _inline(p, metin):
    """**kalın** ve `kod` işaretlerini run'lara çevirir."""
    for parca in re.split(r"(\*\*.+?\*\*|`[^`]+`)", metin):
        if not parca:
            continue
        if parca.startswith("**") and parca.endswith("**"):
            p.add_run(parca[2:-2]).bold = True
        elif parca.startswith("`") and parca.endswith("`"):
            r = p.add_run(parca[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0xA0, 0x30, 0x30)
        else:
            p.add_run(parca)


def _tablo_satiri(satir):
    return [h.strip() for h in satir.strip().strip("|").split("|")]


def _ayirici_mi(satir):
    return bool(re.fullmatch(r"\|[\s:|-]+\|", satir.strip()))


def _tablo_ekle(doc, satirlar):
    basliklar = _tablo_satiri(satirlar[0])
    veri = [_tablo_satiri(s) for s in satirlar[2:]]
    t = doc.add_table(rows=1, cols=len(basliklar))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, b in enumerate(basliklar):
        hucre = t.rows[0].cells[i]
        hucre.text = ""
        p = hucre.paragraphs[0]
        _inline(p, b)
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(9)
        _shade(hucre._tc.get_or_add_tcPr(), BASLIK_ARKA)

    for satir in veri:
        hucreler = t.add_row().cells
        for i, deger in enumerate(satir[:len(basliklar)]):
            p = hucreler[i].paragraphs[0]
            _inline(p, deger)
            for r in p.runs:
                r.font.size = Pt(9)
    doc.add_paragraph()


def _yatay_cizgi(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    alt = OxmlElement("w:bottom")
    alt.set(qn("w:val"), "single")
    alt.set(qn("w:sz"), "6")
    alt.set(qn("w:color"), "BFBFBF")
    bdr.append(alt)
    pPr.append(bdr)


def _icindekiler(doc):
    doc.add_heading("İçindekiler", level=1)
    p = doc.add_paragraph()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), r'TOC \o "1-3" \h \z \u')
    ph = OxmlElement("w:p")
    pr = OxmlElement("w:r")
    tt = OxmlElement("w:t")
    tt.text = "Word'de F9 ile güncelleyin."
    pr.append(tt)
    ph.append(pr)
    fld.append(ph)
    p._p.append(fld)
    doc.add_page_break()


def donustur(md_yolu: Path, docx_yolu: Path):
    satirlar = md_yolu.read_text(encoding="utf-8").splitlines()
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2.0)

    icindekiler_eklendi = False
    i = 0
    while i < len(satirlar):
        s = satirlar[i]
        cikplak = s.strip()

        # kod bloğu
        if cikplak.startswith("```"):
            i += 1
            blok = []
            while i < len(satirlar) and not satirlar[i].strip().startswith("```"):
                blok.append(satirlar[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            _shade(p._p.get_or_add_pPr(), KOD_ARKA)
            r = p.add_run("\n".join(blok))
            r.font.name = "Consolas"
            r.font.size = Pt(8.5)
            continue

        # tablo
        if cikplak.startswith("|") and i + 1 < len(satirlar) and _ayirici_mi(satirlar[i + 1]):
            blok = []
            while i < len(satirlar) and satirlar[i].strip().startswith("|"):
                blok.append(satirlar[i])
                i += 1
            _tablo_ekle(doc, blok)
            continue

        if not cikplak:
            i += 1
            continue

        if cikplak.startswith("---"):
            _yatay_cizgi(doc)
            i += 1
            continue

        # başlıklar
        m = re.match(r"^(#{1,4})\s+(.*)$", cikplak)
        if m:
            seviye = len(m.group(1))
            baslik = m.group(2).replace("**", "")
            h = doc.add_heading(baslik, level=seviye)
            for r in h.runs:
                r.font.color.rgb = ACCENT
            if seviye == 1 and not icindekiler_eklendi:
                _icindekiler(doc)
                icindekiler_eklendi = True
            i += 1
            continue

        # alıntı
        if cikplak.startswith(">"):
            blok = []
            while i < len(satirlar) and satirlar[i].strip().startswith(">"):
                blok.append(satirlar[i].strip().lstrip(">").strip())
                i += 1
            p = doc.add_paragraph()
            _shade(p._p.get_or_add_pPr(), ALINTI_ARKA)
            p.paragraph_format.left_indent = Cm(0.5)
            _inline(p, " ".join(x for x in blok if x))
            continue

        # listeler
        m = re.match(r"^(\s*)[-*]\s+(.*)$", s)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            if len(m.group(1)) >= 2:
                p.paragraph_format.left_indent = Cm(1.5)
            _inline(p, m.group(2))
            i += 1
            continue

        m = re.match(r"^(\s*)\d+\.\s+(.*)$", s)
        if m:
            p = doc.add_paragraph(style="List Number")
            _inline(p, m.group(2))
            i += 1
            continue

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _inline(p, cikplak)
        i += 1

    doc.save(docx_yolu)
    return docx_yolu


if __name__ == "__main__":
    kok = Path(__file__).resolve().parent.parent
    girdi = Path(sys.argv[1]) if len(sys.argv) > 1 else kok / "docs" / "GELISMIS_Sistem_Mimarisi_v3.md"
    if not girdi.is_absolute():
        girdi = kok / girdi
    cikti = Path(sys.argv[2]) if len(sys.argv) > 2 else girdi.with_suffix(".docx")
    if not cikti.is_absolute():
        cikti = kok / cikti
    yol = donustur(girdi, cikti)
    print(f"Yazıldı: {yol}  ({yol.stat().st_size/1024:.0f} KB)")
